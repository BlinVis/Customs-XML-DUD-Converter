"""Conversion of a "perkthim" Excel sheet into an ASYCUDA DUD item XML.

The XML mirrors what Excel produces when the "XML FORMA" template (with its
XML map) is exported: one <Item> per invoice line, empty elements omitted,
tab indentation and CRLF line endings. The old ASYCUDA client imports it.

This module has no GUI code so it can be tested and reused on its own.
"""

from __future__ import annotations

import ast
import datetime as _dt
import operator
import re
from dataclasses import dataclass, field
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

# Logical columns the app needs from the perkthim sheet: (key, required).
FIELDS = [
    ("kodi", False),
    ("perkthim", True),
    ("sasia", True),
    ("totali", True),
    ("pesha_bruto", True),
    ("pesha_neto", True),
    ("origjina", True),
    ("pref", False),
    ("kodi_tarifor", True),
]

DEFAULT_COLUMN_NAMES = {
    "kodi": "Kodi",
    "perkthim": "Perkthim",
    "sasia": "Sasia",
    "totali": "Totali",
    "pesha_bruto": "Pesha bruto",
    "pesha_neto": "Pesha neto",
    "origjina": "Origjina",
    "pref": "Pref",
    "kodi_tarifor": "Kodi tarifor",
}

# Values that were "dragged down" in the template. Editable in the app.
DEFAULT_CONSTANTS = {
    "Number_of_packages": "0",
    "Kind_of_packages_code": "PA",
    "Kind_of_packages_name": "Paketa",
    "National_customs_procedure": "000",
    "Valuation_method_code": "1",
    "Rate_of_adjustement": "1",
    "Currency_code": "EUR",
    "Currency_name": "Nuk ka valutë të huaj",
    "Currency_rate": "1",
}

PRECISION_1_CHOICES = ("00", "90")
EXTENDED_PROCEDURE_CHOICES = ("1000", "4000")

DEFAULT_POSITIVE_INDICATORS = ["Po", "Yes", "Y", "Da", "Ja", "+", "X", "MSA", "True", "1"]
KNOWN_NEGATIVE_INDICATORS = {"jo", "no", "n", "ne", "nein", "-", "false", "0"}

# EU member states (ISO 3166 alpha-2, plus EL which some documents use for Greece).
EU_COUNTRIES = {
    "AT", "BE", "BG", "HR", "CY", "CZ", "DK", "EE", "FI", "FR", "DE", "GR", "EL",
    "HU", "IE", "IT", "LV", "LT", "LU", "MT", "NL", "PL", "PT", "RO", "SK", "SI",
    "ES", "SE",
}

HEADER_SCAN_ROWS = 30


class ConversionError(Exception):
    """A problem that stops the conversion. `key`/`params` feed the translations."""

    def __init__(self, key: str, **params):
        super().__init__(key, params)
        self.key = key
        self.params = params


@dataclass
class Warning_:
    key: str
    params: dict = field(default_factory=dict)


@dataclass
class Options:
    precision_1: str
    extended_procedure: str
    column_names: dict = field(default_factory=lambda: dict(DEFAULT_COLUMN_NAMES))
    positive_indicators: list = field(default_factory=lambda: list(DEFAULT_POSITIVE_INDICATORS))
    constants: dict = field(default_factory=lambda: dict(DEFAULT_CONSTANTS))


@dataclass
class Item:
    row: int
    commodity_code: str
    preference_code: str
    country: str
    description: str
    gross_weight: str
    net_weight: str
    amount: str


@dataclass
class Result:
    xml: str
    items: list
    warnings: list
    sheet_name: str
    header_row: int
    total_amount: float
    total_gross: float
    total_net: float
    indicator_values: dict  # distinct indicator value -> counted as positive?

    @property
    def msa_count(self) -> int:
        return sum(1 for i in self.items if i.preference_code)


# --------------------------------------------------------------------------
# Small helpers


def normalize_header(value) -> str:
    """Header comparison is case-insensitive and ignores extra spaces."""
    if value is None:
        return ""
    return " ".join(str(value).split()).casefold()


def is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def to_number(value):
    """Turn a cell value into a float. Accepts text such as '1.234,56 €'. None if empty."""
    if is_blank(value):
        return None
    if isinstance(value, bool):
        raise ValueError(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[^\d,.\-]", "", str(value))
    if not text or text in "-.,":
        raise ValueError(value)
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif text.count(",") == 1:
        text = text.replace(",", ".")
    elif text.count(",") > 1:
        text = text.replace(",", "")
    return float(text)


def format_number(value: float, decimals: int) -> str:
    """Round and drop trailing zeros, the way Excel writes General numbers."""
    text = f"{round(value, decimals):.{decimals}f}"
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    if text in ("-0", ""):
        text = "0"
    return text


def format_commodity_code(value) -> str:
    if is_blank(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\D", "", str(value))


def output_filename(excel_path, today: _dt.date | None = None) -> str:
    today = today or _dt.date.today()
    return f"DUDI-{Path(excel_path).stem}-{today:%d%m%Y}.xml"


# --------------------------------------------------------------------------
# Formula fallback. Files saved by Excel carry cached results, which are used
# first; files written by other tools may not, so simple arithmetic formulas
# (=E2*D2, =ROUND(E2*D2,2), =SUM(F2:F9)) are evaluated here.

_REF = r"\$?([A-Z]{1,3})\$?(\d+)"
_OPS = {
    ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul,
    ast.Div: operator.truediv, ast.Pow: operator.pow,
    ast.USub: operator.neg, ast.UAdd: operator.pos,
}


class _Sheet:
    def __init__(self, values_ws, formulas_ws):
        self.values = values_ws
        self.formulas = formulas_ws

    def value(self, row: int, col: int, _depth: int = 0):
        cached = self.values.cell(row, col).value
        if cached is not None:
            return cached
        raw = self.formulas.cell(row, col).value
        if isinstance(raw, str) and raw.startswith("="):
            if _depth > 50:
                raise ValueError("formula nesting too deep")
            return self._evaluate(raw[1:], _depth + 1)
        return raw

    def _number_at(self, col_letters, row, depth):
        v = self.value(int(row), column_index_from_string(col_letters), depth)
        n = to_number(v)
        return 0.0 if n is None else n

    def _evaluate(self, expr: str, depth: int):
        expr = expr.upper().replace(";", ",")

        def expand_range(m):
            c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
            i1, i2 = column_index_from_string(c1), column_index_from_string(c2)
            vals = [
                repr(self._number_at(get_column_letter(c), r, depth))
                for r in range(min(r1, r2), max(r1, r2) + 1)
                for c in range(min(i1, i2), max(i1, i2) + 1)
            ]
            return ",".join(vals) or "0"

        expr = re.sub(_REF + ":" + _REF, expand_range, expr)
        expr = re.sub(_REF, lambda m: repr(self._number_at(m.group(1), m.group(2), depth)), expr)
        return _safe_eval(ast.parse(expr, mode="eval").body)


def _safe_eval(node):
    if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
        return float(node.value)
    if isinstance(node, ast.BinOp) and type(node.op) in _OPS:
        return _OPS[type(node.op)](_safe_eval(node.left), _safe_eval(node.right))
    if isinstance(node, ast.UnaryOp) and type(node.op) in _OPS:
        return _OPS[type(node.op)](_safe_eval(node.operand))
    if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
        args = [_safe_eval(a) for a in node.args]
        if node.func.id == "SUM":
            return sum(args)
        if node.func.id == "ROUND" and len(args) == 2:
            return round(args[0], int(args[1]))
    raise ValueError("unsupported formula")


# --------------------------------------------------------------------------
# Reading the workbook


def _load(path):
    path = Path(path)
    if path.suffix.lower() == ".xls":
        raise ConversionError("err_xls")
    try:
        values = openpyxl.load_workbook(path, data_only=True, read_only=False)
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=False)
    except FileNotFoundError:
        raise ConversionError("err_file_missing", path=str(path))
    except Exception as exc:  # corrupt or not an Excel file
        raise ConversionError("err_file_unreadable", path=str(path), detail=str(exc))
    return values, formulas


def list_headers(path) -> list:
    """(sheet name, header row, [header texts]) for each sheet, for showing the user."""
    values, _ = _load(path)
    out = []
    for ws in values.worksheets:
        best_row, best = None, []
        for r in range(1, min(ws.max_row, HEADER_SCAN_ROWS) + 1):
            texts = [str(c.value).strip() for c in ws[r] if isinstance(c.value, str) and c.value.strip()]
            if len(texts) > len(best):
                best_row, best = r, texts
        if best:
            out.append((ws.title, best_row, best))
    return out


def _find_columns(values_wb, column_names):
    """Locate the header row: the first row (any sheet) holding every required name."""
    wanted = {k: normalize_header(column_names.get(k, "")) for k, _ in FIELDS}
    required = [k for k, req in FIELDS if req]
    for k in required:
        if not wanted[k]:
            raise ConversionError("err_column_name_empty", field=k)

    best = None  # (found count, sheet, row, mapping)
    for ws in values_wb.worksheets:
        for r in range(1, min(ws.max_row, HEADER_SCAN_ROWS) + 1):
            headers = {}
            for c in ws[r]:
                key = normalize_header(c.value)
                if key and key not in headers:
                    headers[key] = c.column
            mapping = {k: headers[n] for k, n in wanted.items() if n and n in headers}
            found = sum(1 for k in required if k in mapping)
            if found == len(required):
                missing_optional = [k for k, req in FIELDS if not req and wanted[k] and k not in mapping]
                if missing_optional:
                    raise ConversionError(
                        "err_columns_missing",
                        missing=[column_names[k] for k in missing_optional],
                        sheet=ws.title,
                        found=[str(c.value).strip() for c in ws[r] if not is_blank(c.value)],
                    )
                return ws.title, r, mapping
            if found and (best is None or found > best[0]):
                best = (found, ws, r, mapping)

    if best is None:
        raise ConversionError("err_no_header", names=[column_names[k] for k in required])
    _, ws, r, mapping = best
    raise ConversionError(
        "err_columns_missing",
        missing=[column_names[k] for k in required if k not in mapping],
        sheet=ws.title,
        found=[str(c.value).strip() for c in ws[r] if not is_blank(c.value)],
    )


# --------------------------------------------------------------------------
# Conversion


def convert(excel_path, options: Options) -> Result:
    if options.precision_1 not in PRECISION_1_CHOICES:
        raise ConversionError("err_precision")
    if options.extended_procedure not in EXTENDED_PROCEDURE_CHOICES:
        raise ConversionError("err_extended")

    values_wb, formulas_wb = _load(excel_path)
    sheet_name, header_row, cols = _find_columns(values_wb, options.column_names)
    sheet = _Sheet(values_wb[sheet_name], formulas_wb[sheet_name])
    positives = {normalize_header(v) for v in options.positive_indicators if not is_blank(v)}
    names = options.column_names

    def cell(row, key):
        if key not in cols:
            return None
        try:
            return sheet.value(row, cols[key])
        except (ValueError, SyntaxError, ZeroDivisionError, RecursionError):
            raise ConversionError("err_formula", row=row, column=names[key])

    def number(row, key, warnings):
        raw = cell(row, key)
        try:
            n = to_number(raw)
        except ValueError:
            raise ConversionError("err_not_number", row=row, column=names[key], value=str(raw))
        if n is None:
            warnings.append(Warning_("warn_empty", {"row": row, "column": names[key]}))
        return n

    items, warnings, indicator_values = [], [], {}
    total_amount = total_gross = total_net = 0.0
    if "pref" not in cols:
        warnings.append(Warning_("warn_no_pref_column"))

    max_row = values_wb[sheet_name].max_row
    for r in range(header_row + 1, max_row + 1):
        perkthim = cell(r, "perkthim")
        kodi = cell(r, "kodi")
        if is_blank(perkthim) and is_blank(kodi):
            continue  # empty line or the totals row

        kodi_text = "" if is_blank(kodi) else str(kodi).strip()
        if isinstance(kodi, float) and kodi.is_integer():
            kodi_text = str(int(kodi))
        desc = "" if is_blank(perkthim) else " ".join(str(perkthim).split())
        if not desc:
            warnings.append(Warning_("warn_empty", {"row": r, "column": names["perkthim"]}))
        if kodi_text and not desc.casefold().startswith(kodi_text.casefold()):
            desc = f"{kodi_text} - {desc}" if desc else kodi_text

        qty = number(r, "sasia", warnings)
        desc += f"={format_number(qty or 0, 3)}PCE"

        amount = number(r, "totali", warnings)
        gross = number(r, "pesha_bruto", warnings)
        net = number(r, "pesha_neto", warnings)

        hs = format_commodity_code(cell(r, "kodi_tarifor"))
        if not hs:
            warnings.append(Warning_("warn_empty", {"row": r, "column": names["kodi_tarifor"]}))

        origin_raw = cell(r, "origjina")
        country = "" if is_blank(origin_raw) else str(origin_raw).strip().upper()
        if not country:
            warnings.append(Warning_("warn_empty", {"row": r, "column": names["origjina"]}))

        pref_code = ""
        if "pref" in cols:
            ind = cell(r, "pref")
            ind_key = normalize_header(ind)
            positive = bool(ind_key) and ind_key in positives
            if ind_key:
                indicator_values[str(ind).strip()] = positive
            if positive:
                if country in EU_COUNTRIES:
                    pref_code = "MSA"
                else:
                    warnings.append(Warning_("warn_pref_not_eu", {"row": r, "country": country or "?"}))
            elif ind_key and ind_key not in KNOWN_NEGATIVE_INDICATORS:
                warnings.append(Warning_("warn_unknown_indicator", {"row": r, "value": str(ind).strip()}))

        amount_r = round(amount, 2) if amount is not None else None
        total_amount += amount_r or 0
        total_gross += gross or 0
        total_net += net or 0
        items.append(Item(
            row=r,
            commodity_code=hs,
            preference_code=pref_code,
            country=country,
            description=desc,
            gross_weight="" if gross is None else format_number(gross, 3),
            net_weight="" if net is None else format_number(net, 3),
            amount="" if amount_r is None else format_number(amount_r, 2),
        ))

    if not items:
        raise ConversionError("err_no_rows", sheet=sheet_name)

    return Result(
        xml=build_xml(items, options),
        items=items,
        warnings=warnings,
        sheet_name=sheet_name,
        header_row=header_row,
        total_amount=round(total_amount, 2),
        total_gross=round(total_gross, 3),
        total_net=round(total_net, 3),
        indicator_values=indicator_values,
    )


def build_xml(items, options: Options) -> str:
    c = {k: str(options.constants.get(k, v)).strip() for k, v in DEFAULT_CONSTANTS.items()}
    lines = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<ASYCUDA xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">',
    ]

    def el(depth, tag, value):
        # Like Excel's XML export, an empty cell produces no element at all.
        if value != "":
            lines.append("\t" * depth + f"<{tag}>{escape(value)}</{tag}>")

    def open_(depth, tag):
        lines.append("\t" * depth + f"<{tag}>")

    def close(depth, tag):
        lines.append("\t" * depth + f"</{tag}>")

    for it in items:
        open_(1, "Item")
        open_(2, "Packages")
        el(3, "Number_of_packages", c["Number_of_packages"])
        el(3, "Kind_of_packages_code", c["Kind_of_packages_code"])
        el(3, "Kind_of_packages_name", c["Kind_of_packages_name"])
        close(2, "Packages")
        open_(2, "Tarification")
        open_(3, "HScode")
        el(4, "Commodity_code", it.commodity_code)
        el(4, "Precision_1", options.precision_1)
        close(3, "HScode")
        el(3, "Extended_customs_procedure", options.extended_procedure)
        el(3, "National_customs_procedure", c["National_customs_procedure"])
        el(3, "Valuation_method_code", c["Valuation_method_code"])
        el(3, "Preference_code", it.preference_code)
        close(2, "Tarification")
        open_(2, "Goods_description")
        el(3, "Country_of_origin_code", it.country)
        el(3, "Commercial_Description", it.description)
        close(2, "Goods_description")
        lines.append("\t\t<Previous_doc/>")  # Summary_declaration is always left empty
        open_(2, "Valuation_item")
        open_(3, "Weight_itm")
        el(4, "Gross_weight_itm", it.gross_weight)
        el(4, "Net_weight_itm", it.net_weight)
        close(3, "Weight_itm")
        el(3, "Rate_of_adjustement", c["Rate_of_adjustement"])
        open_(3, "Item_Invoice")
        el(4, "Amount_foreign_currency", it.amount)
        el(4, "Currency_code", c["Currency_code"])
        el(4, "Currency_name", c["Currency_name"])
        el(4, "Currency_rate", c["Currency_rate"])
        close(3, "Item_Invoice")
        close(2, "Valuation_item")
        close(1, "Item")
    lines.append("</ASYCUDA>")
    return "\r\n".join(lines)


def write_xml(xml: str, path) -> None:
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write(xml)
